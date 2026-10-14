#!/usr/bin/env python3
"""Read PFD / mass-balance inputs (PDF, XLSX, DOCX) so their numbers can be transcribed exactly.

PDF inputs may be one file with many sheets, many one-sheet files, or any mix. Commands that take
PDFS accept several files and/or folders (every *.pdf inside); pages are reported as `file.pdf p27`.

Subcommands
  info   PDFS...                          per page: text-layer check, own drawing number, title
  find   PDFS... -t TOKEN [TOKEN ...]     pages whose text contains each token (stream numbers, tags)
  links  PDFS... [--json OUT.json]        sheet index + off-page connectors between sheets: which sheet
         sends which stream to which drawing (FROM/TO), whether that drawing was provided, and
         whether the other sheet shows the same connection back. Use it to follow a flow that
         continues across sheets and to see which drawings are missing from the input.
  text   FILE.pdf PAGE [--raw]            text of one page (pdftotext -layout by default)
  render FILE.pdf PAGES OUTDIR [--dpi N]  render pages to PNG (PAGES like 27 or 3-5 or 6,9,11)
  crop   IMAGE.png OUT.png --box X0 Y0 X1 Y1 [--cols A:B,C:D] [--ref 2000] [--zoom Z]
         Crop a region given in "reference" coordinates (the ~2000 px wide preview the Read tool
         shows). --cols keeps only those x-ranges (e.g. the label column plus the stream columns
         you need) and pastes them side by side, which makes small stream-table digits legible.
  xlsx   FILE.xlsx [--sheet NAME] [--max-rows N] [--max-cols N]
         dump every sheet as tab-separated rows with Excel coordinates (computed values, merged
         cells filled), ready to map onto stream fields.
  docx   FILE.docx                        paragraphs and tables (tab-separated) in document order.

Needs: poppler-utils (pdftoppm, pdftotext, pdfinfo), Pillow, openpyxl, python-docx.
"""
import argparse
import re
import subprocess
import sys
from pathlib import Path


def run(cmd):
    return subprocess.run(cmd, capture_output=True, text=True, errors="replace").stdout


def page_count(pdf):
    m = re.search(r"Pages:\s+(\d+)", run(["pdfinfo", pdf]))
    return int(m.group(1)) if m else 0


def page_text(pdf, page, layout=True):
    cmd = ["pdftotext", "-f", str(page), "-l", str(page)]
    if layout:
        cmd.append("-layout")
    return run(cmd + [pdf, "-"])


def parse_pages(spec):
    out = []
    for part in str(spec).split(","):
        if "-" in part:
            a, b = part.split("-")
            out.extend(range(int(a), int(b) + 1))
        else:
            out.append(int(part))
    return out


def expand_pdfs(items):
    out = []
    for it in items:
        p = Path(it)
        if p.is_dir():
            out.extend(sorted(str(x) for x in p.glob("*.pdf")) + sorted(str(x) for x in p.glob("*.PDF")))
        else:
            out.append(str(p))
    return out


def all_pages(pdfs):
    for f in expand_pdfs(pdfs):
        for p in range(1, page_count(f) + 1):
            yield f, p


def where(f, p):
    return f"{Path(f).name} p{p}"


# Drawing numbers: hyphenated ids that contain a PFD / PID / P&ID part,
# e.g. 123456-00789-9000-PR-PFD-00001, ABC2-9000-PR-PFD-00011, AB-100-PFD-002.
DWG = re.compile(r"[A-Z0-9][A-Z0-9.]*(?:-[A-Z0-9.&]+)*-(?:PFD|PID|P&ID|PDF)-[A-Z0-9]+")
BADGE = re.compile(r"^\d{2,4}[a-z]?(?:,[a-z])*$")  # stream badges: 980, 928a, 404a,b
NOTE_WORDS = ("REFER", "N°", "SEE ", "DOCUMENT", "DETAIL", "LATEST", "REVISION")


def norm_dwg(s):
    s = s.strip().rstrip(".,;:)")
    return re.sub(r"_\d+$", "", s)


def canon_dwg(d):
    """Comparison key that tolerates zero-padding typos (123456-000789-… vs 123456-00789-…)."""
    return "-".join(str(int(x)) if x.isdigit() else x for x in d.split("-"))


def page_lines(f, p):
    """Text lines of one page with positions normalised to 0..1 (x0, y0, x1, y1, text)."""
    h = run(["pdftotext", "-f", str(p), "-l", str(p), "-bbox-layout", f, "-"])
    m = re.search(r'<page width="([\d.]+)" height="([\d.]+)"', h)
    if not m:
        return []
    W, H = float(m.group(1)), float(m.group(2))
    # rotated sheets (/Rotate 90 or 270) report words in the displayed orientation
    xs = [float(v) for v in re.findall(r'xMax="([\d.]+)"', h)]
    ys = [float(v) for v in re.findall(r'yMax="([\d.]+)"', h)]
    if (xs and max(xs) > W * 1.01) or (ys and max(ys) > H * 1.01):
        W, H = H, W
    out = []
    for blk in re.findall(r"<block [^>]*>(.*?)</block>", h, re.S):
        for ln in re.findall(r"<line [^>]*>(.*?)</line>", blk, re.S):
            ws = re.findall(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>', ln)
            if not ws:
                continue
            x0 = min(float(w[0]) for w in ws) / W
            y0 = min(float(w[1]) for w in ws) / H
            x1 = max(float(w[2]) for w in ws) / W
            y1 = max(float(w[3]) for w in ws) / H
            text = " ".join(w[4] for w in ws).replace("&amp;", "&")
            out.append((x0, y0, x1, y1, text))
    return out


def sheet_identity(lines):
    """Own drawing number = a drawing id in the title-block corner (bottom right);
    prefer the one also named in a plotted .dwg path. Returns (own, aliases)."""
    # title-block cells hold the number alone on its line; notes that cite other drawings do not
    corner = [norm_dwg(t) for i, (x0, y0, x1, y1, t) in enumerate(lines)
              if x0 > 0.6 and y0 > 0.83 and DWG.fullmatch(norm_dwg(t).strip())
              and not any(re.match(r"^(FROM|TO)\b", l[4]) for l in lines[i + 1:i + 3])]
    paths = [norm_dwg(d) for *_, t in lines if ".dwg" in t.lower() for d in DWG.findall(t)]
    own = next((d for d in corner if d in paths), None) or (corner[0] if corner else (paths[0] if paths else None))
    aliases = sorted({d for d in corner if d != own})
    return own, aliases


def page_title(f, p):
    t = page_text(f, p, layout=False)
    m = re.search(r"PROCESS FLOW DIAGRAM\s*\n?(.{0,160})", t)
    return (" ".join(m.group(1).split())[:110] if m else ""), len(t.strip())


def connectors(lines, own, aliases):
    """Off-page connectors: a drawing id with a FROM/TO line under it (or sitting at the page edge),
    plus the nearest stream badge on the same row."""
    badges = [(x0, y0, x1, y1, t) for x0, y0, x1, y1, t in lines if all(BADGE.match(w) for w in t.split())]
    found = []
    for i, (x0, y0, x1, y1, t) in enumerate(lines):
        for d in DWG.findall(t):
            d = norm_dwg(d)
            if d == own or d in aliases or ".dwg" in t.lower():
                continue
            nxt = [l[4] for l in lines[i + 1:i + 3]]
            prev = lines[i - 1][4] if i > 0 else ""
            direction = None
            for s in nxt:
                if re.match(r"^(FROM|TO)\b", s.strip()):
                    direction = s.split()[0]
                    break
            note = any(k in t.upper() for k in NOTE_WORDS) or (direction is None and y0 > 0.8)
            if note:
                found.append(dict(kind="note", dwg=d, text=t[:140]))
                continue
            if direction is None:
                if x1 < 0.2:
                    direction = "FROM"
                elif x0 > 0.75:
                    direction = "TO"
                else:
                    found.append(dict(kind="note", dwg=d, text=t[:140]))
                    continue
            cy = (y0 + y1) / 2
            best = None
            for bx0, by0, bx1, by1, bt in badges:
                dy = abs((by0 + by1) / 2 - cy)
                dx = bx0 - x1 if direction == "FROM" else x0 - bx1
                if dy < 0.012 and -0.02 < dx < 0.25 and (best is None or dx < best[0]):
                    best = (dx, bt)
            label = prev if not DWG.search(prev) and not re.match(r"^(FROM|TO)\b", prev) and not BADGE.match(prev) else ""
            detail = next((s for s in nxt if re.match(r"^(FROM|TO)\b", s.strip())), "")
            found.append(dict(kind="flow", dir=direction, dwg=d, stream=best[1] if best else "",
                              label=label[:60], detail=detail[:90], y=round(cy, 3)))
    return found


def cmd_info(a):
    for f in expand_pdfs(a.files):
        print(f"=== {f}")
        print(run(["pdfinfo", f]).strip())
        for p in range(1, page_count(f) + 1):
            title, chars = page_title(f, p)
            own, aliases = sheet_identity(page_lines(f, p))
            flag = "text" if chars > 200 else "IMAGE-ONLY (needs visual reading/OCR)"
            print(f"p{p:>3} [{flag}, {chars} chars] {own or '(drawing no. not found)'}"
                  + (f" = {', '.join(aliases)}" if aliases else "") + (f"  {title}" if title else ""))
        print()


def cmd_find(a):
    texts = {(f, p): page_text(f, p, layout=False) for f, p in all_pages(a.files)}
    for tok in a.tokens:
        pat = re.compile(r"(?<![\w.-])" + re.escape(tok) + r"(?![\w.-])")
        hits = [where(f, p) for (f, p), t in texts.items() if pat.search(t)]
        print(f"{tok}: {', '.join(hits) if hits else '-'}")


def cmd_links(a):
    import json
    sheets, index = [], {}
    for f, p in all_pages(a.files):
        lines = page_lines(f, p)
        own, aliases = sheet_identity(lines)
        title, chars = page_title(f, p)
        sh = dict(file=Path(f).name, page=p, dwg=own, aliases=aliases, title=title,
                  text_layer=chars > 200, connectors=connectors(lines, own, aliases))
        sheets.append(sh)
        for d in [own] + aliases:
            if d:
                index.setdefault(d, []).append(sh)
                index.setdefault(canon_dwg(d), []).append(sh) if canon_dwg(d) != d else None

    def sheet_of(d):
        hit = index.get(d) or index.get(canon_dwg(d))
        return hit[0] if hit else None

    def loc(d):
        hit = sheet_of(d)
        return where(hit["file"], hit["page"]) if hit else None

    print("== Sheets")
    for sh in sheets:
        print(f"{where(sh['file'], sh['page']):<28} {sh['dwg'] or '(no drawing no.)':<34} {sh['title']}"
              + ("" if sh["text_layer"] else "  [IMAGE-ONLY: connectors not detected]"))
    dups = {d: v for d, v in index.items() if len(v) > 1 and "-" in d and d == (v[0]["dwg"] or "")}
    for d, v in dups.items():
        print(f"WARN  {d} appears as own number on {', '.join(where(s['file'], s['page']) for s in v)}")

    print("\n== Connections between sheets (direction as seen from the sheet in column 1)")
    missing, flows = {}, []
    for sh in sheets:
        for c in sh["connectors"]:
            if c["kind"] != "flow":
                continue
            tgt = loc(c["dwg"])
            back = "?"
            if tgt:
                other = sheet_of(c["dwg"])
                want = "TO" if c["dir"] == "FROM" else "FROM"
                mine = {canon_dwg(x) for x in [sh["dwg"]] + sh["aliases"] if x}
                ok = any(o["kind"] == "flow" and o["dir"] == want and canon_dwg(o["dwg"]) in mine
                         and (not c["stream"] or not o["stream"] or o["stream"] == c["stream"])
                         for o in other["connectors"])
                back = "matched" if ok else "no return connector"
                if c["dwg"] not in index:
                    back += " (dwg no. typo in source)"
            else:
                missing.setdefault(c["dwg"], []).append(where(sh["file"], sh["page"]))
            flows.append(dict(sheet=sh["dwg"], at=where(sh["file"], sh["page"]), **c, target=tgt, check=back))
            print(f"{where(sh['file'], sh['page']):<20} {c['dir']:<4} {c['stream'] or '-':<7} {c['dwg']:<34} "
                  f"{(tgt or 'NOT PROVIDED'):<20} {back:<20} {c['label']} {c['detail']}")
    if missing:
        print("\n== Referenced drawings not in the input")
        for d, at in sorted(missing.items()):
            print(f"{d:<34} referenced by {', '.join(sorted(set(at)))}")
    if a.json:
        Path(a.json).write_text(json.dumps(dict(sheets=sheets, flows=flows, missing=missing), ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"\nwrote {a.json}")


def cmd_text(a):
    sys.stdout.write(page_text(a.file, a.page, layout=not a.raw))


def cmd_render(a):
    out = Path(a.outdir)
    out.mkdir(parents=True, exist_ok=True)
    for p in parse_pages(a.pages):
        prefix = out / f"p{p:02d}"
        subprocess.run(["pdftoppm", "-f", str(p), "-l", str(p), "-r", str(a.dpi), "-png", "-singlefile", a.file, str(prefix)],
                       stderr=subprocess.DEVNULL, check=True)
        print(f"{prefix}.png")


def cmd_crop(a):
    from PIL import Image
    im = Image.open(a.image)
    W, H = im.size
    s = W / a.ref
    x0, y0, x1, y1 = a.box
    ranges = [(x0, x1)]
    if a.cols:
        ranges = [tuple(float(v) for v in c.split(":")) for c in a.cols.split(",")]
    parts = [im.crop((int(cx0 * s), int(y0 * s), int(cx1 * s), int(y1 * s))) for cx0, cx1 in ranges]
    w = sum(p.size[0] for p in parts)
    h = max(p.size[1] for p in parts)
    out = Image.new("RGB", (w, h), "white")
    x = 0
    for p in parts:
        out.paste(p, (x, 0))
        x += p.size[0]
    if a.zoom and a.zoom != 1:
        out = out.resize((int(w * a.zoom), int(h * a.zoom)))
    out.save(a.out)
    print(f"{a.out} {out.size[0]}x{out.size[1]} (source scale {s:.2f}px per ref unit)")


def cmd_xlsx(a):
    import openpyxl
    wb = openpyxl.load_workbook(a.file, data_only=True)
    wbf = openpyxl.load_workbook(a.file, data_only=False)  # formulas, shown when no cached value exists
    sheets = [wb[a.sheet]] if a.sheet else wb.worksheets
    for ws in sheets:
        wsf = wbf[ws.title]
        print(f"=== sheet: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
        merged = {}
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged[(r, c)] = v
        for r in range(1, min(ws.max_row, a.max_rows) + 1):
            vals = []
            for c in range(1, min(ws.max_column, a.max_cols) + 1):
                v = ws.cell(r, c).value
                if v is None:
                    v = merged.get((r, c))
                if v is None:
                    f = wsf.cell(r, c).value
                    if isinstance(f, str) and f.startswith("="):
                        v = f"[no cached value: {f}]"
                vals.append("" if v is None else str(v).replace("\n", " ").strip())
            while vals and vals[-1] == "":
                vals.pop()
            if any(vals):
                print(f"R{r}\t" + "\t".join(vals))
        print()


def cmd_docx(a):
    import docx
    d = docx.Document(a.file)
    body = d.element.body
    ti = 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in child.iter() if n.tag.endswith("}t")).strip()
            if t:
                print(t)
        elif tag == "tbl":
            tbl = d.tables[ti]
            ti += 1
            print(f"--- table {ti} ({len(tbl.rows)} rows)")
            for row in tbl.rows:
                cells = [c.text.replace("\n", " ").strip() for c in row.cells]
                print("\t".join(cells))
            print("--- end table")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("info"); p.add_argument("files", nargs="+"); p.set_defaults(fn=cmd_info)
    p = sub.add_parser("find"); p.add_argument("files", nargs="+"); p.add_argument("-t", "--tokens", nargs="+", required=True); p.set_defaults(fn=cmd_find)
    p = sub.add_parser("links"); p.add_argument("files", nargs="+"); p.add_argument("--json"); p.set_defaults(fn=cmd_links)
    p = sub.add_parser("text"); p.add_argument("file"); p.add_argument("page", type=int); p.add_argument("--raw", action="store_true"); p.set_defaults(fn=cmd_text)
    p = sub.add_parser("render"); p.add_argument("file"); p.add_argument("pages"); p.add_argument("outdir"); p.add_argument("--dpi", type=int, default=250); p.set_defaults(fn=cmd_render)
    p = sub.add_parser("crop"); p.add_argument("image"); p.add_argument("out")
    p.add_argument("--box", nargs=4, type=float, required=True, metavar=("X0", "Y0", "X1", "Y1"))
    p.add_argument("--cols", help="x-ranges to keep, e.g. 50:215,420:600")
    p.add_argument("--ref", type=float, default=2000, help="width of the reference coordinate space (default 2000)")
    p.add_argument("--zoom", type=float, default=1.0)
    p.set_defaults(fn=cmd_crop)
    p = sub.add_parser("xlsx"); p.add_argument("file"); p.add_argument("--sheet"); p.add_argument("--max-rows", type=int, default=400); p.add_argument("--max-cols", type=int, default=80); p.set_defaults(fn=cmd_xlsx)
    p = sub.add_parser("docx"); p.add_argument("file"); p.set_defaults(fn=cmd_docx)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
