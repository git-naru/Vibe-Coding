"""The app's Excel downloads (app.html `downloadEquipment`, `downloadEstimate`) with the same sheets and cells.

Column widths are set from the app's character widths; Excel's rendering of them may differ slightly from
SheetJS's. The '만든 때' (created) cell holds the time of writing, as in the app.
"""
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter

from . import js
from .cost import HOW, eq_key
from .equipment import row_flags
from .estimate import fmt_cap


def ko_now(now=None):
    """new Date().toLocaleString('ko-KR'): 2026. 10. 15. 오후 3:04:05"""
    t = now or datetime.now()
    h = t.hour % 12 or 12
    return f"{t.year}. {t.month}. {t.day}. {'오전' if t.hour < 12 else '오후'} {h}:{t.minute:02d}:{t.second:02d}"


def _r2(x):
    return js.js_round(x * 100) / 100


def _sheet(wb, name, aoa, widths=None, autofilter=False):
    ws = wb.create_sheet(name)
    for r, row in enumerate(aoa, start=1):
        for c, v in enumerate(row, start=1):
            if v is None:
                continue
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"   # text, as SheetJS writes it (openpyxl would take it for a formula)
    for j, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if autofilter and aoa and aoa[0]:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(aoa[0]))}{len(aoa)}"
    return ws


def equipment_workbook(rows, refs_missing, area_names, cost=None, costs=None, mb_title="", inputs="", vision="", now=None):
    """The Equipment List workbook. cost = the record's cost state, costs = compute_costs(cost, rows);
    inputs = the '입력' line (e.g. "a.pdf (3장)"), vision = "1/3장 실행"."""
    mb = any(r.get("mb") for r in rows)
    K, C = costs, cost
    cu = f" ({C['unit']})" if C and C.get("unit") else ""

    def list_rows(items):
        lst = (C or {}).get("list")
        return ", ".join(js.jstr(lst["items"][i]["row"]) if lst and 0 <= i < len(lst["items"]) else "" for i in (items or []))

    def cost_cols(r):
        if not K:
            return []
        b = K["byKey"].get(eq_key(r))
        if not b:
            return ["", "", ""]
        return [_r2(b["cost"]), js.jstr(HOW.get(b["how"])) + (f" (1행을 {b['split']}기에 나눔)" if b.get("split") else ""), list_rows(b["items"])]

    head = ["No.", "Area", "Tag No.", "Equipment Name (도면 원문)", "종류", "운전 구분", *([f"설비비{cu}", "설비비 매칭", "설비비 목록 행"] if K else []),
            "비고 (도면 표기)", "Nom. Vol. (m³)", "PFD 도면 번호", "PDF 파일", "PDF 쪽", "도면 구역 (행-열)",
            *(["값 출처 탭", "입구 Stream", "출구 Stream", "기준 Stream", "유체", "일 유량 (t/d)", "Nominal 평균 (t/h)", "Design 최대 (t/h)", "운전 온도 (°C)", "운전 압력 (bar)", "밀도 (t/m³)", "고형분 (%)"] if mb else []),
            "확인 필요"]
    aoa = [head]
    for i, r in enumerate(rows):
        m = r.get("mb")
        aoa.append([i + 1, r["area"], r["tag"], r["name"], r["kind"], r["status"], *cost_cols(r), r["note"], r["vol"], r["dwg"], r["file"], r["page"], r["zone"],
                    *(([m["tab"], m["in"], m["out"], m["basis"], m["fluid"], m["dm"], m["nm"], m["dx"], m["T"], m["P"], m["rho"], m["sol"]] if m else [""] * 12) if mb else []),
                    row_flags(r)])

    def wch(h):
        return 44 if js.test("Name", h) else 32 if js.test("도면 번호|확인", h) else 22 if js.test("파일|유체", h) else max(8, min(16, js.length(h) + 2))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Equipment List", aoa, [wch(h) for h in head], autofilter=True)
    # summary by area × kind (values; the list sheet is the source)
    counts = {}
    for r in rows:
        counts[r["kind"]] = counts.get(r["kind"], 0) + 1
    kinds = [k for k, _ in js.sort(js.entries(counts), lambda a, b: b[1] - a[1])][:9]
    areas = sorted(js.unique(r["area"] for r in rows))
    summary = [["Area", "Area 이름", "전체", *kinds, "기타", "운전", "예비", "향후", "미확인 (Tag 없음)", *([f"설비비 합계{cu}"] if K else [])]]
    for a in areas:
        rs = [r for r in rows if r["area"] == a]
        k = [sum(1 for r in rs if r["kind"] == x) for x in kinds]
        summary.append([a, area_names.get(a) or "(입력에 이 Area 도면 없음)", len(rs), *k, len(rs) - js.add(k),
                        sum(1 for r in rs if r["status"] == "운전"), sum(1 for r in rs if r["status"].startswith("예비")),
                        sum(1 for r in rs if r["status"].startswith("향후")), sum(1 for r in rs if r["status"] == "미확인"),
                        *([_r2(js.add(((K["byKey"].get(eq_key(r)) or {}).get("cost") or 0) for r in rs))] if K else [])])
    summary.append(["합계", "", *[js.add(row[j + 2] for row in summary[1:]) for j in range(len(summary[0]) - 2)]])
    _sheet(wb, "Summary", summary, [44 if j == 1 else 10 for j in range(len(summary[0]))])
    _sheet(wb, "도면 미포함 설비", [["Tag No.", "언급된 시트", "판단"], *[[x["tag"], x["sheets"], x["why"]] for x in refs_missing]], [15, 40, 70])
    if K and C.get("list"):   # the manual list, row by row, with the PFD equipment each row went to
        def key_label(k):
            r = rows[K["keyOf"].index(k)] if k in K["keyOf"] else None
            return ("(Tag 없음) " + r["name"] if r.get("untagged") else r["tag"]) if r else k
        cm = [["목록 행", "Tag (원문)", "설비명", "수량", f"금액{cu}", "매칭된 PFD Tag", "매칭 방법"]]
        for it in C["list"]["items"]:
            ks = K["assign"].get(js.jstr(it["i"])) or []
            cm.append([it["row"], it["tagRaw"], it["name"], "" if it["qty"] is None else it["qty"], "" if it["cost"] is None else it["cost"],
                       ", ".join(key_label(k) for k in ks), (HOW.get(K["how"].get(js.jstr(it["i"]))) or "") if ks else "PFD에 없음"])
        cm += [[], ["수동 목록 합계", "", "", "", _r2(K["total"])], ["PFD 설비에 매칭", "", "", "", _r2(K["total"] - K["outside"])], ["PFD에 없음", "", "", "", _r2(K["outside"])]]
        _sheet(wb, "설비비 매칭", cm, [8, 26, 44, 7, 16, 34, 18])
    basis = [["항목", "기준"], ["입력", inputs],
             ["설비 판별", '설비 이름 아래에 단독으로 적힌 Tag(NNNN-XX-NNN, 앞에 WF- 같은 접두가 붙을 수 있음)를 설비로 본다. 괄호·FROM/TO 문구·주석 속 Tag는 다른 시트 설비를 가리키는 참조라 "도면 미포함 설비" 탭에서 따로 본다.'],
             ["이름", '도면 원문 그대로. 이름 하나 아래 Tag가 여럿이면 "PUMP 1/2", "A/B", "2@8" 표기를 Tag 순서대로 나눠 붙였다.'],
             ["운전 구분", "(STAND-BY)·(SPARE)는 묶음의 마지막 Tag(또는 그 표시가 옆에 붙은 Tag)에, (FUTURE)는 묶음 전체에. 표시가 없으면 운전. Tag 없이 이름만 있는 설비 후보는 \"미확인\"."],
             ["도면 구역", '도면 테두리 격자(위쪽 열 번호, 왼쪽 행 문자)로 Tag 글자(미확인은 이름 글자) 위치를 "행-열"로 적었다.'],
             ["미확인(Tag 없음)", '도면 글자에서 설비 이름만 있고 Tag가 없는 것을 찾은 후보다. 이름이 "(그림)"으로 시작하는 행은 "그림으로 설비 찾기"에서 Claude가 도면 이미지를 보고 찾은, 글자 없이 기호로만 그려진 설비 후보이며 확인 필요 열에 신뢰도를 적었다. 트럭·창고·피트·방류벽은 넣지 않았다.'],
             ["그림 대조", vision]]
    if K:
        lst = C.get("list")
        basis.append(["설비비", (f"수동 Equipment List {lst['file']} ({lst['sheet']} 시트)를 Tag로 먼저 맞추고(번호 자릿수·범위 표기 001/002·001~003 정리), Tag가 없는 행은 이름이 비슷한 설비, 그다음 Claude 제안, 사람이 고른 매칭 순으로 붙였다. 목록 1행이 여러 기를 가리키면 금액을 기수로 나눴다. " if lst else "")
                      + f"표에서 직접 입력한 금액은 \"직접 입력\"으로 표시했다. 단위: {C.get('unit') or '미지정'}."])
    basis.append(["물질수지 값", f"흐름도 탭({mb_title})의 Stream 값. 메인 흐름 탭을 먼저 보고, 없으면 다른 탭 값을 쓴다. 펌프·팬·압축기·보일러는 출구, 그 밖은 입구 Stream 기준이며 여러 개면 유량 합산, 온도·압력은 범위." if mb
                  else "흐름도를 생성하거나 열면 그 계통 설비에 Stream 값이 채워진다."])
    basis.append(["만든 때", ko_now(now)])
    _sheet(wb, "작성 기준", basis, [18, 110])
    return wb


_METHOD = {"expo": "용량지수", "unit": "실적 단가", "skip": "제외", "none": "실적 없음"}
_PART = {"eq": "설비", "pkg": "패키지", "common": "공통"}


def estimate_workbook(calc, rates, project_name, now=None):
    """The estimate basis workbook (calc = est_compute(...))."""
    C = calc
    cu, ru = "/".join(C["units"]), rates.get("currency") or ""

    def r2(x):
        return "" if x is None else _r2(x)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "요약", [["항목", "금액", "단위"], ["설비비", r2(C["eqTotal"]), cu], ["공사비", r2(C["wkTotal"]), ru],
                       ["투자비 (설비비 + 공사비)", r2(C["eqTotal"] + C["wkTotal"]), cu if cu == ru else f"{cu} + {ru}"], [],
                       ["프로젝트", project_name], ["비교 실적", ", ".join(js.jstr(d["p"].get("name")) for d in C["refs"])], ["만든 때", ko_now(now)]], [26, 18, 14])
    _sheet(wb, "설비비", [["Tag No.", "설비 이름", "종류", "용량", "용량 단위", "유사 실적 설비", "실적 프로젝트", "실적 용량", "실적 설비비", "용량 지수 n", "용량비", "용량지수 산출", "실적 단가 (중앙값)", "단가 표본 수", "단가 산출", "적용 방법", f"설비비 ({cu})", "직접 입력", "확인"],
                         *[[e["r"]["tag"], e["r"]["name"], e["r"]["kind"], r2(e["cap"]["v"]) if e["cap"] else "", e["cap"]["u"] if e["cap"] else "",
                            e["ref"]["tag"] if e["ref"] else "", e["ref"]["projName"] if e["ref"] else "", fmt_cap(e["ref"]["cap"]) if e["ref"] else "",
                            r2(e["ref"]["cost"]) if e["ref"] else "", e["n"], r2(e["ratio"]), r2(e["expo"]), r2(e["up"]), e["upN"], r2(e["unitEst"]),
                            _METHOD.get(e["method"]), r2(e["val"]), "예" if e["over"] else "", "; ".join(e["flags"])] for e in C["eqRes"]]],
           [14, 36, 12, 10, 8, 14, 18, 14, 14, 8, 8, 14, 14, 8, 14, 10, 14, 8, 40])
    _sheet(wb, "공사비", [["공종", "단위", "설비 기준 물량", "패키지 기준 물량", "공통 물량", "적용 물량", "직접 입력", f"단가 ({ru})", f"공사비 ({ru})"],
                         *[[w["work"], w["unit"] or w["rateUnit"], r2(w["eq"]), r2(w["pkg"]), r2(w["common"]), r2(w["q"]), "예" if w["qOver"] else "", r2(w["rate"]), r2(w["amount"])] for w in C["works"]]],
           [18, 8, 14, 14, 12, 12, 8, 12, 14])
    _sheet(wb, "공사물량 근거", [["공종", "구분", "물량", "근거"], *[[w["work"], _PART[l["part"]], r2(l["q"]), l["why"]] for w in C["works"] for l in w["lines"]]], [18, 8, 12, 90])
    q = js.jstr(rates.get("qexp"))
    _sheet(wb, "산출 기준", [["항목", "내용"],
                          ["설비비", '같은 종류의 실적 설비 중 용량(같은 단위)이 가장 가까운 설비를 골라 실적 설비비 × (용량비)^n. 용량을 비교할 수 없으면 실적 금액 그대로. "실적 단가"는 같은 종류·같은 단위 실적의 설비비/용량 중앙값 × 용량.'],
                          ["용량", '탱크·반응조·시크너 등은 Nominal 용적(m³), 그 밖은 흐름도 Stream의 Design(없으면 Nominal) 유량(t/h). 화면에서 고친 값은 "✎"로 표시.'],
                          ["공사물량", f"설비 기준: 실적에서 설비 종류·공종별 물량 / Σ(그 종류 설비비^{q})를 학습해 이 프로젝트의 같은 종류 Σ(설비비^{q})에 곱한다(실적이 여럿이면 평균). 패키지 기준: 이름(같은 이름 우선, '패키지'·'설비' 같은 일반어 제외)이나 설비 종류 구성이 가장 비슷한 실적 패키지 물량 × (패키지 설비비 비)^{q}. 공통: 실적의 (공통 물량 / 설비비 합계) 평균 × 이 프로젝트 설비비 합계."],
                          ["공사비", "적용 물량 × 회사 공사 단가 (공종 이름으로 연결)."],
                          *[[f"지수 · {'기본' if k == '_기본' else k}", v] for k, v in js.entries(rates["exps"])], ["물량 지수", rates.get("qexp")]], [22, 110])
    return wb
