"""Workbook reading: pfdcore's reader against SheetJS 0.18.5 (raw values, formatted text, Stream tables, CSV)."""
import openpyxl
import pytest
from conftest import assert_same, jsq, run_app

from pfdcore import cost, ssf, streams, xlsx

FORMATS = ["General", "0", "0.0", "0.00", "0.000", "#,##0", "#,##0.0", "#,##0.00", "0%", "0.0%", "0.00%", '#,##0.0" t/h"',
           '0.00 "bar"', "#,##0;(#,##0)", "#,##0.00;[Red]-#,##0.00", "#", "#.##", "00000", "#,##0,", '"$"#,##0.00', "@"]
VALUES = [0, 1, -1, 7, 0.5, 2.5, -2.5, 1562.5, 0.1 + 0.2, 1 / 3, 2 / 3, 12345.678, -12345.678, 123456789012.5, 1e21, -0.000012345,
          0.00045, 99.995, 1.005, 2.675, 1000000, 47.9, 0.256]


@pytest.fixture(scope="module")
def stream_book(tmp_path_factory):
    p = tmp_path_factory.mktemp("x") / "streams.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MB"
    ws["B2"] = "MADE-UP MASS BALANCE"
    ws.append([])
    rows = [["", "Stream", "Unit", "101", "102", "105 a"], ["", "Description", "", "Feed brine", "Pump discharge", "Soft brine"],
            ["", "Balance flow", "", "", "", ""], ["", "Daily mass flow", "t/d", 1200, 1180.5, 1150.25],
            ["", "Nominal flow", "", "", "", ""], ["", "Mass flow", "t/h", 50, 49.2, 47.9], ["", "Volumetric flow", "m3/h", 42.37, 41.69, 40.94],
            ["", "Maximum design flow", "", "", "", ""], ["", "Mass flow", "t/h", 55, 54.1, 52.7],
            ["", "Other values", "", "", "", ""], ["", "Temperature", "°C", 25, 25, 60], ["", "Pressure", "bar g", 1, 4.5, 3.8],
            ["", "Density", "t/m3", 1.18, 1.18, 1.172], ["", "pH", "", 7.2, 7.2, 8.1], ["", "Li", "ppm", 350, 350, 349.5],
            ["", "Chemical composition liquid", "", "", "", ""], ["", "NaCl", "%", 22.1, 22.1, 22.3], ["", "CaCl2", "%", 0.4, 0.4, "-"]]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=4):
        for c in row[3:]:
            if isinstance(c.value, float):
                c.number_format = "#,##0.0"
    ws2 = wb.create_sheet("Formats")
    for i, f in enumerate(FORMATS, start=1):
        for j, v in enumerate(VALUES, start=1):
            c = ws2.cell(row=i, column=j, value=v)
            c.number_format = f
    ws2.cell(row=len(FORMATS) + 1, column=1, value="=1+1")
    ws2.cell(row=len(FORMATS) + 1, column=2, value=True)
    ws2.cell(row=len(FORMATS) + 1, column=3, value="  text with spaces  ")
    wb.create_sheet("Empty")
    wb.save(p)
    return p


def test_raw_and_text(stream_book):
    app = run_app(f"const wb = workbook({jsq(stream_book)}); return wb.SheetNames.map(n => [n, XLSX.utils.sheet_to_json(wb.Sheets[n], {{header: 1, raw: true, defval: ''}}), XLSX.utils.sheet_to_json(wb.Sheets[n], {{header: 1, raw: false, defval: ''}})]);")
    wb = xlsx.read_workbook(stream_book)
    assert [a[0] for a in app] == wb.sheet_names
    ssf.unsupported.clear()
    for name, raw, text in app:
        assert_same(raw, wb.aoa(name, raw=True), f"{name} raw")
        assert_same(text, wb.aoa(name, raw=False), f"{name} text")
    assert not ssf.unsupported, f"formats fell back to General: {ssf.unsupported}"


def test_stream_tables(stream_book):
    app = run_app(f"return parseXlsxTables(workbook({jsq(stream_book)})).map(t => ({{...t, mapped: mapRows(t.rows, t.ids, t.colValues)}}));")
    py = streams.parse_xlsx_tables(xlsx.read_workbook(stream_book))
    for t in py:
        t["mapped"] = streams.map_rows(t["rows"], t["ids"], t["colValues"])
    assert_same(app, py, "stream tables")
    assert py and py[0]["mapped"]["102"]["v"]["dm"] == "1,180.5" and py[0]["desc"]["101"] == "Feed brine"


def test_csv(tmp_path):
    p = tmp_path / "cost.csv"
    p.write_text('Tag No.,Description,Q\'ty,Amount (USD)\n1000-TK-001,FEED TANK,1,"120,000"\n1000-PU-001/002,"PUMP, FEED",2,30000.5\n'
                 ',SAMPLE COOLER,1,4000\n1000-HX-003,HEATER,,TRUE\n,TOTAL,,154000.5\n', encoding="utf-8")
    app = run_app(f"const wb = workbook({jsq(p)}); return {{rows: XLSX.utils.sheet_to_json(wb.Sheets[wb.SheetNames[0]], {{header: 1, raw: true, defval: ''}}), list: parseCostWorkbook(wb)}};")
    wb = xlsx.read_workbook(p)
    assert_same(app["rows"], wb.aoa("Sheet1", raw=True), "CSV rows")
    assert_same(app["list"], cost.parse_cost_workbook(wb), "CSV cost list")
