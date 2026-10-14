"""Made-up inputs for the parity tests: PFD sheets (PDF), an equipment-cost list, quantity and rate lists (XLSX).

Everything here is fictitious — drawing numbers (ABC-…), tags, names, streams and costs are invented. The PDF
writer is hand-rolled so the tests control every text operator pdf.js has to handle.
"""
import zlib
from pathlib import Path

import openpyxl

A3 = (1190.55, 841.89)


def _esc(s):
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


class PdfWriter:
    def __init__(self):
        self.objs = []

    def add(self, body):
        self.objs.append(body)
        return len(self.objs)

    def stream(self, data, extra="", compress=False):
        if isinstance(data, str):
            data = data.encode("latin-1")
        if compress:
            data = zlib.compress(data)
            extra += " /Filter /FlateDecode"
        return self.add(b"<< /Length %d %s >>\nstream\n" % (len(data), extra.encode()) + data + b"\nendstream")

    def save(self, path, root):
        out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
        offs = []
        for i, body in enumerate(self.objs, start=1):
            offs.append(len(out))
            if isinstance(body, str):
                body = body.encode("latin-1")
            out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
        x = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(self.objs) + 1)
        for o in offs:
            out += b"%010d 00000 n \n" % o
        out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(self.objs) + 1, root, x)
        Path(path).write_bytes(bytes(out))


ARIAL_W = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278] + [556] * 10 + [278, 278, 584, 584, 584, 556, 1015] + \
    [667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611] + \
    [278, 278, 278, 469, 556, 333] + [556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500]
HANGUL = "급수탱크펌프보일러응축수"


def _fonts(w):
    f1 = w.add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    desc = w.add("<< /Type /FontDescriptor /FontName /ArialMT /Flags 32 /FontBBox [-665 -325 2000 1006] /ItalicAngle 0 /Ascent 905 /Descent -212 /CapHeight 716 /StemV 80 /MissingWidth 250 >>")
    f2 = w.add(f"<< /Type /Font /Subtype /TrueType /BaseFont /ArialMT /FirstChar 32 /LastChar {31 + len(ARIAL_W)} /Widths [{' '.join(map(str, ARIAL_W))}] /Encoding /WinAnsiEncoding /FontDescriptor {desc} 0 R >>")
    cmap = "/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n" \
           "/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n1 begincodespacerange\n<0000> <FFFF>\nendcodespacerange\n" \
           f"{len(HANGUL)} beginbfchar\n" + "".join(f"<{i + 1:04X}> <{ord(c):04X}>\n" for i, c in enumerate(HANGUL)) + \
           "endbfchar\nendcmap\nCMapName currentdict /CMap defineresource pop\nend\nend"
    tu = w.stream(cmap)
    kdesc = w.add("<< /Type /FontDescriptor /FontName /MadeUpGothic /Flags 4 /FontBBox [0 -200 1000 900] /ItalicAngle 0 /Ascent 900 /Descent -200 /CapHeight 700 /StemV 80 >>")
    kcid = w.add(f"<< /Type /Font /Subtype /CIDFontType2 /BaseFont /MadeUpGothic /CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> "
                 f"/FontDescriptor {kdesc} 0 R /DW 1000 /W [1 [920 920 900] 4 12 880] >>")
    f3 = w.add(f"<< /Type /Font /Subtype /Type0 /BaseFont /MadeUpGothic /Encoding /Identity-H /DescendantFonts [{kcid} 0 R] /ToUnicode {tu} 0 R >>")
    f4 = w.add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    return f"/F1 {f1} 0 R /F2 {f2} 0 R /F3 {f3} 0 R /F4 {f4} 0 R"


def _hangul(s):
    return "<" + "".join(f"{HANGUL.index(c) + 1:04X}" for c in s) + ">"


class Page:
    """Content stream builder in page points (origin bottom left)."""

    def __init__(self, W, H):
        self.W, self.H, self.ops = W, H, []

    def text(self, x, y, s, font="F1", size=7, extra=""):
        self.ops.append(f"BT /{font} {size} Tf {extra} 1 0 0 1 {x:.2f} {y:.2f} Tm ({_esc(s)}) Tj ET")

    def at(self, fx, fy, s, **kw):
        """Text at page fractions (fy from the top, at the text's top edge as the app sees it)."""
        size = kw.get("size", 7)
        self.text(fx * self.W, self.H - fy * self.H - size, s, **kw)

    def raw(self, op):
        self.ops.append(op)

    def data(self):
        return "\n".join(self.ops)


def pfd_sheet(W, H, area, area_name, dwg, title2, labels, connectors, table=None, untagged=(), notes=()):
    p = Page(W, H)
    for i in range(1, 9):   # border grid: column numbers on top, row letters on the left
        p.at((i - 0.5) / 8, 0.005, str(i), size=6)
    for i, ch in enumerate("ABCDEF"):
        p.at(0.004, (i + 0.5) / 6, ch, size=6)
    for fx, fy, lines in labels:   # name lines above the tag line(s), remarks under
        y = fy
        for ln in lines:
            p.at(fx, y, ln)
            y += 0.0105
    for fx, fy, name in untagged:
        p.at(fx, fy, name)
    for fx, fy, d, direction, badge, label in connectors:
        p.at(fx, fy - 0.0105, label)
        p.at(fx, fy, d, size=6)
        p.at(fx, fy + 0.012, f"{direction} {label}", size=5)
        p.at(fx + (0.10 if direction == "FROM" else -0.03), fy, badge, size=6)
    for fx, fy, s in notes:
        p.at(fx, fy, s, size=6)
    if table:
        x0, y0, ids, rows = table
        p.at(x0, y0, "Stream", size=6)
        p.at(x0 + 0.12, y0, "Unit", size=6)
        for k, sid in enumerate(ids):
            p.at(x0 + 0.16 + 0.05 * k, y0, sid, size=6)
        for r, (label, unit, vals) in enumerate(rows):
            yy = y0 + 0.013 * (r + 1)
            p.at(x0, yy, label, size=6)
            p.at(x0 + 0.12, yy, unit, size=6)
            for k, v in enumerate(vals):
                p.at(x0 + 0.16 + 0.05 * k, yy, v, size=6)
    # title block
    p.at(0.80, 0.865, f"{area} - {area_name}", size=8)
    p.at(0.80, 0.885, "PROCESS FLOW DIAGRAM", font="F4", size=8)
    p.at(0.80, 0.905, title2, size=7)
    p.at(0.80, 0.93, dwg, size=8)
    p.at(0.62, 0.975, f"C:\\PROJECTS\\MADEUP\\{dwg}.dwg", size=5)
    return p


def pfd_pdf(path):
    """Two made-up PFD sheets plus a rotated sheet and a sheet with a scanned (image) stream table."""
    w = PdfWriter()
    fonts = _fonts(w)
    W, H = A3
    s1 = pfd_sheet(W, H, "1000", "FEED PREPARATION", "ABC-1000-PR-PFD-00001", "BRINE FEED AND SOFTENING",
                   labels=[(0.20, 0.20, ["BRINE FEED", "TANK", "1000-TK-001", "NOM. VOL: 25 m3"]),
                           (0.35, 0.20, ["BRINE FEED PUMPS 1/2", "1000-PU-001", "1000-PU-002", "(STAND-BY)"]),
                           (0.50, 0.20, ["SOFTENER", "FEED HEATER", "1000-HX-003"]),
                           (0.65, 0.20, ["LIME SLURRY", "TANK", "1000-TK-004", "1000-AG-004", "AGITATOR"]),
                           (0.20, 0.45, ["CARTRIDGE FILTERS", "A/B", "1000-FL-005", "1000-FL-006"]),
                           (0.35, 0.45, ["BLOWERS 2@4", "WF-1000-FA-002", "WF-1000-FA-003", "WF-1000-FA-004"]),
                           (0.50, 0.45, ["DOSING PUMP", "1000-PU-007", "(SPARE)"]),
                           (0.65, 0.45, ["2 SLUDGE POND", "1000-ML-008", "(FUTURE)"])],
                   connectors=[(0.03, 0.62, "ABC-0900-PR-PFD-00009", "FROM", "100", "RAW BRINE"),
                               (0.88, 0.62, "ABC-2000-PR-PFD-00002", "TO", "105", "SOFT BRINE")],
                   table=(0.10, 0.70, ["101", "102", "105"],
                          [("Balance flow", "", ["", "", ""]), ("Daily mass flow", "t/d", ["1,200", "1,180", "1,150"]),
                           ("Nominal flow", "", ["", "", ""]), ("Mass flow", "t/h", ["50.0", "49.2", "47.9"]),
                           ("Other values", "", ["", "", ""]), ("Temperature", "°C", ["25", "25", "60"]),
                           ("Pressure", "bar", ["1.0", "4.5", "3.8"]), ("Density", "t/m3", ["1.18", "1.18", "1.17"])]),
                   untagged=[(0.80, 0.30, "SAMPLE COOLER"), (0.80, 0.40, "TRUCK LOADING")],
                   notes=[(0.05, 0.80, "NOTE 1. ALL TANKS ARE TO BE VENTED TO THE SCRUBBER."), (0.05, 0.81, "SEE 1000-PU-010 / 012 FOR SPARE PUMPS")])
    # text-operator coverage on sheet 1 (all fictitious)
    s1.raw("BT /F2 7 Tf 1 0 0 1 300 120 Tm [(KERNED) -120 (LABEL) 40 (S)] TJ ET")
    s1.raw("BT /F2 7 Tf 2 Tc 1 0 0 1 300 110 Tm (SPACED OUT) Tj 0 Tc ET")
    s1.raw("BT /F2 7 Tf 3 Tw 1 0 0 1 300 100 Tm (WORD SPACING HERE) Tj 0 Tw ET")
    s1.raw("BT /F1 7 Tf 80 Tz 1 0 0 1 300 90 Tm (NARROW 1000-PU-099) Tj 100 Tz ET")
    s1.raw("BT /F1 7 Tf 1 0 0 1 300 80 Tm (FIRST) Tj 20 0 Td (SECOND) Tj 2 Ts (RAISED) Tj 0 Ts ET")
    s1.raw("BT /F1 7 Tf 1 0 0 1 300 70 Tm (FAR) Tj 120 0 Td (APART) Tj ET")
    s1.raw("q 0 1 -1 0 1100 300 cm BT /F1 6 Tf 1 0 0 1 0 0 Tm (ROTATED LABEL 1000-TK-011) Tj ET Q")
    s1.raw("BT /F1 7 Tf 1 0 0 1 300 60 Tm 9 TL (LINE ONE) Tj T* (LINE TWO) Tj (LINE 3) ' ET")
    s1.raw(f"BT /F3 7 Tf 1 0 0 1 500 120 Tm {_hangul('급수탱크')} Tj ET")
    s1.raw("q 1.5 0 0 1.5 600 60 cm BT /F1 5 Tf 0 0 Td (SCALED TEXT) Tj ET Q")
    s1.raw("/Fm1 Do")
    s1.raw("BT /F1 7 Tf 1 0 0 1 420 50 Tm (A  B   C) Tj ET")
    form = w.stream("BT /F1 6 Tf 1 0 0 1 10 10 Tm (FORM TEXT 1000-CV-020) Tj ET", extra=f"/Type /XObject /Subtype /Form /BBox [0 0 200 50] /Matrix [1 0 0 1 700 40] /Resources << /Font << {fonts} >> >>")
    c1 = w.stream(s1.data(), compress=True)
    xo1 = f"/XObject << /Fm1 {form} 0 R >>"

    s2 = pfd_sheet(W, H, "2000", "SOFTENING", "ABC-2000-PR-PFD-00002", "IX SOFTENING AND STORAGE",
                   labels=[(0.20, 0.25, ["SOFT BRINE", "TANK", "2000-TK-001"]),
                           (0.35, 0.25, ["IX COLUMNS", "A/B/C", "2000-IX-002", "2000-IX-003", "2000-IX-004"]),
                           (0.50, 0.25, ["TRANSFER PUMPS", "1/2", "2000-PU-005", "2000-PU-006"]),
                           (0.65, 0.25, ["MISNAMED PUMPS 1/2/3", "2000-PU-007", "2000-PU-008"]),
                           (0.20, 0.50, ["1000-PU-002", "SPARE PUMP DRAWN HERE"])],
                   connectors=[(0.03, 0.62, "ABC-1000-PR-PFD-00001", "FROM", "105", "SOFT BRINE"),
                               (0.88, 0.62, "ABC-3000-PR-PFD-00003", "TO", "201", "TO STORAGE")],
                   untagged=[(0.80, 0.30, "REGENERATION SKID")],
                   notes=[(0.05, 0.80, "REFER TO ABC-0900-PR-PFD-00009 FOR DETAILS.")])
    s2.raw("q 600 0 0 300 400 250 cm /Im1 Do Q")     # a scanned stream table (large image)
    s2.raw("q 20 0 0 20 100 100 cm BI /W 2 /H 2 /CS /G /BPC 8 ID \x00\xff\xff\x00 EI Q")   # small inline image
    img = w.stream(b"\x00\x80\xff\x40", extra="/Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceGray /BitsPerComponent 8")
    c2 = w.stream(s2.data())
    xo2 = f"/XObject << /Im1 {img} 0 R >>"

    # sheet 3: portrait page shown rotated (/Rotate 90), drawing number only in the .dwg path
    W3, H3 = H, W
    s3 = Page(W3, H3)
    s3.raw(f"BT /F1 8 Tf 0 1 -1 0 {W3 - 100:.2f} 300 Tm (3000 - STORAGE) Tj ET")
    s3.raw(f"BT /F4 8 Tf 0 1 -1 0 {W3 - 112:.2f} 300 Tm (PROCESS FLOW DIAGRAM) Tj ET")
    s3.raw(f"BT /F1 7 Tf 0 1 -1 0 {W3 - 124:.2f} 300 Tm (PRODUCT STORAGE) Tj ET")
    s3.raw(f"BT /F1 5 Tf 0 1 -1 0 {W3 - 20:.2f} 700 Tm (C:\\\\PROJECTS\\\\MADEUP\\\\ABC-3000-PR-PFD-00003.dwg) Tj ET")
    s3.raw("BT /F1 7 Tf 0 1 -1 0 400 200 Tm (PRODUCT TANK) Tj ET")
    s3.raw("BT /F1 7 Tf 0 1 -1 0 409 200 Tm (3000-TK-001) Tj ET")
    s3.raw("BT /F1 7 Tf 0 1 -1 0 400 400 Tm (LOADING PUMP) Tj ET")
    s3.raw("BT /F1 7 Tf 0 1 -1 0 409 400 Tm (3000-PU-002) Tj ET")
    c3 = w.stream(s3.data())

    kids = []
    pages_id = len(w.objs) + 4
    for c, xo, (pw, ph), extra in [(c1, xo1, (W, H), ""), (c2, xo2, (W, H), ""), (c3, "", (W3, H3), "/Rotate 90")]:
        kids.append(w.add(f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 {pw:.2f} {ph:.2f}] {extra} /Contents {c} 0 R "
                          f"/Resources << /Font << {fonts} >> {xo} >> >>"))
    w.add(f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>")
    assert len(w.objs) == pages_id
    root = w.add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>")
    w.save(path, root)
    return path


def cost_xlsx(path):
    """A made-up manual equipment list with costs (USD)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EQ COST"
    ws["A1"] = "MADE-UP PLANT — EQUIPMENT COST"
    ws.append([])
    ws.append(["No.", "Tag No.", "Description", "Q'ty", "Unit Price (USD)", "Amount (USD)"])
    rows = [
        (1, "1000-TK-001", "BRINE FEED TANK", 1, 120000, 120000),
        (2, "1000-PU-001/002", "BRINE FEED PUMP", 2, 15000, 30000),
        (3, "1000-HX-3", "SOFTENER FEED HEATER", 1, 45000, None),
        (4, "1000-TK-004", "LIME SLURRY TANK", 1, 60000, 60000),
        (5, "1000-AG-004", "LIME SLURRY TANK AGITATOR", 1, 18000, 18000),
        (6, "1000-FL-005~006", "CARTRIDGE FILTER", 2, 9000, 18000),
        (7, "", "SAMPLE COOLER", 1, 4000, 4000),
        (8, "1000-PU-099", "SPARE PARTS PUMP", 1, 2500, 2500),
        (9, "2000-IX-002/003/004", "IX COLUMN", 3, 80000, 240000),
        (10, "2000-PU-005", "TRANSFER PUMP", 1, 12000, 12000),
        (11, "", "INSTALLATION SUPERVISION", 1, 30000, 30000),
        (12, "WF-1000-FA-002", "BLOWER 2", 1, 7000, 7000),
        (None, "", "TOTAL", None, None, 541500),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def qty_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quantities"
    ws.append(["대상", "공종", "물량", "단위"])
    for r in [("1000-PU-001~002", "배관", 120, "m"), ("1000-TK-001", "기초 콘크리트", 30, "m³"), ("1000-TK-004", "기초 콘크리트", 20, "m³"),
              ("Feed 패키지", "철골", 12, "t"), ("공통", "토목", 500, "m³"), ("", "전기 케이블", 2000, "m"),
              ("9999-XX-001", "도장", 50, "m²"), ("합계", "", None, "")]:
        ws.append(list(r))
    wb.save(path)
    return path


def rates_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["공종", "단위", "단가 (USD)"])
    for r in [("배관", "m", 150), ("기초 콘크리트", "m³", 400), ("철골", "t", 3200), ("토목", "m³", 60), ("전기 케이블", "m", 25)]:
        ws.append(list(r))
    wb.save(path)
    return path


def cjk_pdf(path):
    """Labels in fonts pdf.js can only read with its CMap files (not embedded, text made up): Korean in an
    Adobe-Korea1 CID font without /ToUnicode and in a UniKS-UCS2-H font (the app carries those CMaps), and
    Japanese in an Adobe-Japan1 CID font without /ToUnicode (it does not)."""
    from pdfminer.cmapdb import CMapDB
    cid_of, jcid_of = {}, {}
    for cid, u in CMapDB.get_unicode_map("Adobe-Korea1", False).cid2unichr.items():
        cid_of.setdefault(u, cid)
    for cid, u in CMapDB.get_unicode_map("Adobe-Japan1", False).cid2unichr.items():
        jcid_of.setdefault(u, cid)
    w = PdfWriter()
    desc = w.add("<< /Type /FontDescriptor /FontName /MadeUpMyungjo /Flags 6 /FontBBox [0 -200 1000 900] /ItalicAngle 0 /Ascent 900 /Descent -200 /CapHeight 700 /StemV 80 >>")
    k1 = w.add(f"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /MadeUpMyungjo /CIDSystemInfo << /Registry (Adobe) /Ordering (Korea1) /Supplement 2 >> /FontDescriptor {desc} 0 R /DW 1000 >>")
    f1 = w.add(f"<< /Type /Font /Subtype /Type0 /BaseFont /MadeUpMyungjo /Encoding /Identity-H /DescendantFonts [{k1} 0 R] >>")
    k2 = w.add(f"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /MadeUpMyungjo /CIDSystemInfo << /Registry (Adobe) /Ordering (Korea1) /Supplement 2 >> /FontDescriptor {desc} 0 R /DW 900 >>")
    f2 = w.add(f"<< /Type /Font /Subtype /Type0 /BaseFont /MadeUpMyungjo /Encoding /UniKS-UCS2-H /DescendantFonts [{k2} 0 R] >>")
    f3 = w.add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    jdesc = w.add("<< /Type /FontDescriptor /FontName /MadeUpMincho /Flags 6 /FontBBox [0 -200 1000 900] /ItalicAngle 0 /Ascent 900 /Descent -200 /CapHeight 700 /StemV 80 >>")
    j1 = w.add(f"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /MadeUpMincho /CIDSystemInfo << /Registry (Adobe) /Ordering (Japan1) /Supplement 4 >> /FontDescriptor {jdesc} 0 R /DW 1000 >>")
    f4 = w.add(f"<< /Type /Font /Subtype /Type0 /BaseFont /MadeUpMincho /Encoding /Identity-H /DescendantFonts [{j1} 0 R] >>")
    hexcid = lambda s: "<" + "".join(f"{cid_of[c]:04X}" for c in s) + ">"
    ucs2 = lambda s: "<" + "".join(f"{ord(c):04X}" for c in s) + ">"
    ops = [f"BT /F3 7 Tf 1 0 0 1 100 500 Tm (2000-TK-001) Tj ET",
           f"BT /F1 7 Tf 1 0 0 1 100 512 Tm {hexcid('급수탱크')} Tj ET",
           f"BT /F2 7 Tf 1 0 0 1 100 450 Tm {ucs2('보일러')} Tj ET",
           f"BT /F3 7 Tf 1 0 0 1 100 440 Tm (2000-BO-002) Tj ET",
           "BT /F4 7 Tf 1 0 0 1 100 380 Tm <" + "".join(f"{jcid_of[c]:04X}" for c in "ポンプ") + "> Tj ET"]
    c = w.stream("\n".join(ops))
    pages_id = len(w.objs) + 2
    page = w.add(f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 842 595] /Contents {c} 0 R /Resources << /Font << /F1 {f1} 0 R /F2 {f2} 0 R /F3 {f3} 0 R /F4 {f4} 0 R >> >> >>")
    w.add(f"<< /Type /Pages /Kids [{page} 0 R] /Count 1 >>")
    root = w.add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>")
    w.save(path, root)
    return path


def layout_pdf(path, title=True):
    """A made-up equipment arrangement drawing: tags written in full, as ranges, short ('PU-007' without the area),
    off the list, and in the title block. With title=False only the file name says it is a Layout."""
    w = PdfWriter()
    fonts = _fonts(w)
    W, H = A3
    p = Page(W, H)
    for fx, fy, s in [(0.20, 0.20, "1000-TK-001"), (0.30, 0.22, "1000-PU-001/002"), (0.40, 0.25, "PU-007"), (0.45, 0.25, "TK-004"),
                      (0.50, 0.30, "2000-IX-002~004"), (0.55, 0.35, "3000-PU-099"), (0.60, 0.40, "4000-ZZ-001"), (0.62, 0.45, "HX-003 / FL-005"),
                      (0.25, 0.50, "BRINE FEED TANK"), (0.70, 0.86, "5000-TK-001"), (0.30, 0.60, "PU-001")]:
        p.at(fx, fy, s)
    if title:
        p.at(0.80, 0.885, "GENERAL ARRANGEMENT", font="F4", size=8)
        p.at(0.80, 0.905, "FEED AREA PLOT", size=7)
    else:
        p.at(0.80, 0.885, "AREA 1000", size=8)
    c = w.stream(p.data())
    pid = len(w.objs) + 2
    pg = w.add(f"<< /Type /Page /Parent {pid} 0 R /MediaBox [0 0 {W:.2f} {H:.2f}] /Contents {c} 0 R /Resources << /Font << {fonts} >> >> >>")
    w.add(f"<< /Type /Pages /Kids [{pg} 0 R] /Count 1 >>")
    root = w.add(f"<< /Type /Catalog /Pages {pid} 0 R >>")
    w.save(path, root)
    return path
